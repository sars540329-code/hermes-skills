#!/usr/bin/env python3
"""
generate_word_report.py — 策略分析 Word 報告產生器

吃兩個 JSON（config + metrics）+ 圖表資料夾，產出標準 10 節 Word。
文字需要 agent 依 word-delivery.md 的風格自填（透過 --narrative narrative.json 傳入）。

輸入：
  --config config.json       回測配置（策略名稱、標的、資金、部位規則、手續費…）
  --metrics metrics.json     calc_metrics.py 的輸出
  --narrative narrative.json 每節的文字內容（見範例 schema）
  --charts charts/           plot_charts.py / rolling_perf.py 的 PNG
  --out 報告.docx

narrative.json schema：
{
  "rating": "🟢 綠燈",
  "one_liner": "...",
  "key_numbers": ["年賺...", "最慘...", ...],
  "s2_analogy": "...",
  "s3_backtest_setup": "...",
  "s3_results": "...",
  "s4_risk": "...",
  "s5_worst_case": "...",
  "s6_regimes": "...",
  "s7_sizing": "...",
  "s8_alerts": "...",
  "s9_limitations": "...",
  "s10_conclusion": "..."
}

用法：
  python generate_word_report.py \
      --config config.json --metrics metrics.json \
      --narrative narrative.json --charts charts/ --out 策略報告.docx
"""
from __future__ import annotations
import argparse
import json
from pathlib import Path
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8")   # Windows 終端機預設 cp950，印中文／emoji 會爆
except Exception:
    pass

try:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    sys.exit("這支需要 python-docx（Word 報告是選配交付物）。\n"
             "安裝：pip install python-docx\n"
             "不想裝也沒關係 —— 其他腳本與 HTML 報告都不需要它。")


SECTION_TITLES = [
    "一、一頁摘要",
    "二、策略在做什麼",
    "三、過去表現如何",
    "四、風險會有多大",
    "五、最壞情境是什麼",
    "六、在不同市況下的表現",
    "七、要用多少錢、做多大部位",
    "八、什麼時候要停用",
    "九、可能失準的地方",
    "十、結論與下一步",
]


def _add_picture_if_exists(doc: Document, path: Path, width_inches: float = 6.0):
    if path and path.exists():
        doc.add_picture(str(path), width=Inches(width_inches))


def _add_paragraph(doc: Document, text: str, bold: bool = False, size: int = 11):
    if not text:
        return
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)


def _add_table(doc: Document, table_spec: dict):
    """
    table_spec = {"title": "...", "headers": [...], "rows": [[...], ...]}
    """
    if not table_spec:
        return
    title = table_spec.get("title")
    headers = table_spec.get("headers", [])
    rows = table_spec.get("rows", [])
    if title:
        p = doc.add_paragraph()
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(11)
    if not headers:
        return
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Light Grid Accent 1"
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = str(h)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(10)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            if j >= len(headers):
                break
            cell = t.rows[i].cells[j]
            cell.text = str(val)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(10)
    doc.add_paragraph()  # spacing


def _add_section(doc: Document, narrative: dict, key_text: str, key_tables: str = None):
    """Add paragraph(s) + optional list of tables referenced by key_tables."""
    _add_paragraph(doc, narrative.get(key_text, ""))
    if key_tables:
        for spec in narrative.get(key_tables, []) or []:
            _add_table(doc, spec)


def build(config: dict, metrics: dict, narrative: dict, charts: Path, out: Path):
    doc = Document()
    title = doc.add_heading(f'{config.get("strategy_name", "策略")} 投資策略完整分析報告', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # § 1 一頁摘要
    doc.add_heading(SECTION_TITLES[0], level=1)
    _add_paragraph(doc, f'📊 策略評級：{narrative.get("rating", "")}', bold=True, size=13)
    _add_paragraph(doc, narrative.get("one_liner", ""))
    _add_paragraph(doc, "關鍵數字：", bold=True)
    for n in narrative.get("key_numbers", []):
        doc.add_paragraph(n, style="List Bullet")

    # § 2 策略在做什麼
    doc.add_heading(SECTION_TITLES[1], level=1)
    _add_section(doc, narrative, "s2_analogy", "s2_tables")

    # § 3 過去表現
    doc.add_heading(SECTION_TITLES[2], level=1)
    _add_paragraph(doc, "先說清楚這次回測是怎麼跑的：", bold=True)
    _add_paragraph(doc, narrative.get("s3_backtest_setup", ""))
    for spec in narrative.get("s3_tables_setup", []) or []:
        _add_table(doc, spec)
    _add_paragraph(doc, narrative.get("s3_results", ""))
    for spec in narrative.get("s3_tables_results", []) or []:
        _add_table(doc, spec)
    _add_picture_if_exists(doc, charts / "equity_curve.png")
    _add_picture_if_exists(doc, charts / "monthly_heatmap.png")

    # § 4 風險
    doc.add_heading(SECTION_TITLES[3], level=1)
    _add_section(doc, narrative, "s4_risk", "s4_tables")

    # § 5 最壞情境
    doc.add_heading(SECTION_TITLES[4], level=1)
    _add_section(doc, narrative, "s5_worst_case", "s5_tables")
    _add_picture_if_exists(doc, charts / "drawdown.png")
    _add_picture_if_exists(doc, charts / "rolling_sharpe.png")

    # § 6 不同市況
    doc.add_heading(SECTION_TITLES[5], level=1)
    _add_section(doc, narrative, "s6_regimes", "s6_tables")
    _add_picture_if_exists(doc, charts / "is_oos.png")
    _add_picture_if_exists(doc, charts / "wfa_equity.png")

    # § 7 配置
    doc.add_heading(SECTION_TITLES[6], level=1)
    _add_section(doc, narrative, "s7_sizing", "s7_tables")
    _add_picture_if_exists(doc, charts / "mc_distribution.png")

    # § 8 停用條件
    doc.add_heading(SECTION_TITLES[7], level=1)
    _add_section(doc, narrative, "s8_alerts", "s8_tables")

    # § 9 限制
    doc.add_heading(SECTION_TITLES[8], level=1)
    _add_section(doc, narrative, "s9_limitations", "s9_tables")

    # § 10 結論
    doc.add_heading(SECTION_TITLES[9], level=1)
    _add_section(doc, narrative, "s10_conclusion", "s10_tables")

    # 附錄：原始指標表（calc_metrics.py 可能輸出字串 "inf"，不能直接用 :.3f）
    def _fmt(v, spec=".3f", pct=False):
        if isinstance(v, (int, float)):
            return f"{v*100:{spec}}%" if pct else f"{v:{spec}}"
        return str(v) if v is not None else "—"

    doc.add_page_break()
    doc.add_heading("附錄：原始績效指標", level=1)
    lines = [
        f'CAGR            {_fmt(metrics.get("cagr"), ".2f", pct=True)}',
        f'Sharpe Ratio    {_fmt(metrics.get("sharpe"))}',
        f'Sortino Ratio   {_fmt(metrics.get("sortino"))}',
        f'Calmar Ratio    {_fmt(metrics.get("calmar"))}',
        f'Max Drawdown    {_fmt(metrics.get("max_drawdown", {}).get("value"), ".2f", pct=True)}',
        f'Win Rate        {_fmt(metrics.get("win_rate"), ".2f", pct=True)}' if "win_rate" in metrics else "",
        f'Profit Factor   {_fmt(metrics.get("profit_factor"))}' if "profit_factor" in metrics else "",
        f'Trades          {metrics.get("trades_total", "")}' if "trades_total" in metrics else "",
    ]
    for line in lines:
        if line:
            p = doc.add_paragraph(line)
            p.runs[0].font.name = "Consolas"
            p.runs[0].font.size = Pt(10)

    sc = metrics.get("scorecard")
    if sc:
        _add_table(doc, {
            "title": f'五指標燈號（加總 {sc.get("summary", "")}）',
            "headers": ["指標", "數值", "燈號"],
            "rows": [[r["metric"], _fmt(r["value"]), r["light"]] for r in sc.get("rows", [])],
        })

    _add_paragraph(doc, "本報告為量化方法學習與回測研究，不構成任何投資建議；歷史回測不代表未來績效。",
                   bold=True, size=10)

    doc.save(out)
    return out


def main():
    ap = argparse.ArgumentParser(description="生成策略 Word 交付報告")
    ap.add_argument("--config", type=Path, required=True)
    ap.add_argument("--metrics", type=Path, required=True)
    ap.add_argument("--narrative", type=Path, required=True)
    ap.add_argument("--charts", type=Path, default=Path("charts/"))
    ap.add_argument("--out", type=Path, required=True)
    args = ap.parse_args()

    config = json.loads(args.config.read_text(encoding="utf-8-sig"))
    metrics = json.loads(args.metrics.read_text(encoding="utf-8-sig"))
    narrative = json.loads(args.narrative.read_text(encoding="utf-8-sig"))

    out = build(config, metrics, narrative, args.charts, args.out)
    print(f"[ok] wrote {out}")


if __name__ == "__main__":
    main()
